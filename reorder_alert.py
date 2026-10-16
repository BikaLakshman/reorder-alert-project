"""
Automated Inventory Reorder Alert
Reads inventory.csv, works out how many days of stock are left for each item,
flags items that need reordering, and writes a formatted Excel report plus a chart.

Run:  python reorder_alert.py
"""
import pandas as pd
import matplotlib.pyplot as plt
from datetime import date

INPUT_FILE = "inventory.csv"
OUTPUT_FILE = f"reorder_report_{date.today()}.xlsx"
SAFETY_DAYS = 3        # extra buffer days of stock we want to keep
TARGET_WEEKS = 4       # how many weeks of sales a restock order should cover


def load_data(path):
    df = pd.read_csv(path)
    required = {"item_id", "item_name", "category", "stock_on_hand",
                "avg_weekly_sales", "lead_time_days"}
    missing = required - set(df.columns)
    if missing:
        raise ValueError(f"Missing columns in {path}: {missing}")
    return df


def analyze(df):
    df = df.copy()
    df["daily_sales"] = df["avg_weekly_sales"] / 7
    # days until we run out at the current sales rate
    df["days_of_stock"] = (df["stock_on_hand"] / df["daily_sales"]).round(1)
    # reorder point = stock needed to cover the supplier lead time plus a safety buffer
    df["reorder_point"] = (df["daily_sales"] * (df["lead_time_days"] + SAFETY_DAYS)).round(0).astype(int)

    def status(row):
        if row["stock_on_hand"] == 0:
            return "OUT OF STOCK"
        if row["days_of_stock"] <= row["lead_time_days"]:
            return "CRITICAL"   # will run out before a new order could arrive
        if row["stock_on_hand"] <= row["reorder_point"]:
            return "REORDER"
        return "OK"

    df["status"] = df.apply(status, axis=1)
    # suggested order: enough to reach TARGET_WEEKS of sales, plus cover lead time demand
    target_stock = df["avg_weekly_sales"] * TARGET_WEEKS + df["daily_sales"] * df["lead_time_days"]
    df["suggested_order_qty"] = (target_stock - df["stock_on_hand"]).clip(lower=0).round(0).astype(int)
    df.loc[df["status"] == "OK", "suggested_order_qty"] = 0
    order = {"OUT OF STOCK": 0, "CRITICAL": 1, "REORDER": 2, "OK": 3}
    return df.sort_values(by=["status", "days_of_stock"],
                          key=lambda c: c.map(order) if c.name == "status" else c)


def make_chart(df, path="status_chart.png"):
    counts = df["status"].value_counts().reindex(["OUT OF STOCK", "CRITICAL", "REORDER", "OK"], fill_value=0)
    colors = ["#c0392b", "#e67e22", "#f1c40f", "#27ae60"]
    plt.figure(figsize=(6, 4))
    plt.bar(counts.index, counts.values, color=colors)
    plt.title("Items by Inventory Status")
    plt.ylabel("Number of items")
    plt.tight_layout()
    plt.savefig(path, dpi=150)
    plt.close()


def write_excel(df, path):
    summary = (df.groupby("category")
                 .agg(items=("item_id", "count"),
                      needs_action=("status", lambda s: (s != "OK").sum()))
                 .reset_index())
    flagged = df[df["status"] != "OK"]

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        flagged.to_excel(writer, sheet_name="Action Needed", index=False)
        df.to_excel(writer, sheet_name="All Items", index=False)
        summary.to_excel(writer, sheet_name="By Category", index=False)

        fills = {"OUT OF STOCK": "F5B7B1", "CRITICAL": "F8C471", "REORDER": "F9E79F"}
        from openpyxl.styles import PatternFill, Font
        for sheet in ("Action Needed", "All Items"):
            ws = writer.sheets[sheet]
            status_col = [c.value for c in ws[1]].index("status") + 1
            for row in ws.iter_rows(min_row=2):
                color = fills.get(row[status_col - 1].value)
                if color:
                    for cell in row:
                        cell.fill = PatternFill("solid", fgColor=color)
        for ws in writer.sheets.values():
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for col in ws.columns:   # auto-fit column widths
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[col[0].column_letter].width = width + 3
            ws.freeze_panes = "A2"


def main():
    df = analyze(load_data(INPUT_FILE))
    write_excel(df, OUTPUT_FILE)
    make_chart(df)
    counts = df["status"].value_counts()
    print(f"Report saved: {OUTPUT_FILE}")
    print(counts.to_string())
    print(f"\n{(df['status'] != 'OK').sum()} of {len(df)} items need attention.")


if __name__ == "__main__":
    main()
